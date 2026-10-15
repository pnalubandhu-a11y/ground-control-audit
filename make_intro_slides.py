"""Generate intro-slides.pptx — 2 context slides shown before the live demo.

Run:  python3 make_intro_slides.py   (requires python-pptx)
Edit the SLIDES data below and re-run to regenerate.
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

NAVY  = RGBColor(0x0B, 0x14, 0x2E)
PANEL = RGBColor(0x14, 0x20, 0x44)
CYAN  = RGBColor(0x3A, 0xD0, 0xE8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DIM   = RGBColor(0xAE, 0xB9, 0xD0)
AMBER = RGBColor(0xFF, 0xC1, 0x07)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def _box(slide, l, t, w, h, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    return tf


def _run(p, text, size, color, bold=False, italic=False):
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.color.rgb = color
    r.font.bold = bold
    r.font.italic = italic
    r.font.name = "Arial"
    return r


def add_slide(eyebrow, title, bullets, takeaway, takeaway_color, notes):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = NAVY

    # left accent bar
    bar = s.shapes.add_shape(1, Inches(0), Inches(0), Inches(0.18), Inches(7.5))
    bar.fill.solid(); bar.fill.fore_color.rgb = CYAN; bar.line.fill.background()

    # eyebrow
    tf = _box(s, 0.9, 0.55, 11.5, 0.6)
    _run(tf.paragraphs[0], eyebrow, 16, CYAN, bold=True)

    # title
    tf = _box(s, 0.9, 1.15, 11.6, 1.9)
    _run(tf.paragraphs[0], title, 36, WHITE, bold=True)

    # bullets
    tf = _box(s, 0.95, 3.15, 11.5, 2.9)
    for i, b in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(14)
        _run(p, "•  ", 22, CYAN, bold=True)
        _run(p, b, 22, WHITE)

    # takeaway strip
    strip = s.shapes.add_shape(1, Inches(0.9), Inches(6.35), Inches(11.5), Inches(0.72))
    strip.fill.solid(); strip.fill.fore_color.rgb = PANEL; strip.line.fill.background()
    tf = strip.text_frame; tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    _run(p, takeaway, 19, takeaway_color, bold=True)

    s.notes_slide.notes_text_frame.text = notes
    return s


add_slide(
    eyebrow="GROUND CONTROL  -  CoreWeave Moonshot 2026",
    title="When a GPU cluster slows down, nobody knows WHY fast enough",
    bullets=[
        "A stalled 1,000-GPU training run burns thousands of dollars an hour - sitting idle.",
        "Finding the cause means an engineer hand-correlating nvidia-smi, ibstat, NCCL logs and counters across dozens of nodes.",
        "The culprit may be hardware (a bad GPU or cable) or one wrong setting - and both just look like \"the cluster is slow.\"",
    ],
    takeaway="Today, diagnosis is manual, slow, and expert-bound - hours to days.",
    takeaway_color=AMBER,
    notes=(
        "~25s. Set the stakes before the terminal. GPU time is the most expensive "
        "thing we sell; when a run stalls, the clock is running. The painful part "
        "isn't fixing it - it's FINDING it, by hand, across many nodes, and the "
        "cause could be hardware or a single config setting. Both look identical "
        "from the outside: 'it's just slow.' Hold the problem here, then: 'so we "
        "built this.'"
    ),
)

add_slide(
    eyebrow="THE SOLUTION  -  what you're about to see",
    title="Ground Control: one command, a named root cause + fix in 30 seconds",
    bullets=[
        "audit runs 40+ server, network and GPU checks - then a verdict engine correlates the signals into ONE named cause with the fix.",
        "One tool, three audiences: internal (full detail), customer (plain-English, self-serve, redacted), support (ticket bundle).",
        "Deterministic, runs on real CoreWeave Blackwell hardware, and uses zero AI tokens at runtime.",
    ],
    takeaway="Next, live: a healthy node, a failing GPU, and a ~100x NCCL slowdown - each named and explained.",
    takeaway_color=CYAN,
    notes=(
        "~25s. This is the pivot into the demo. The differentiator is the verdict "
        "engine: it doesn't just list red/green checks, it CORRELATES them into one "
        "named root cause with the action. Same tool serves ops, customers, and "
        "support. It's deterministic and runs on real hardware - no AI cost to run. "
        "Then launch the terminal: 'let me show you.'"
    ),
)

prs.save("intro-slides.pptx")
print("wrote intro-slides.pptx  (%d slides)" % len(prs.slides._sldIdLst))
